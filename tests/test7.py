"""Word export from the app.

Drives the browser to build a .docx, saves it, then opens it with python-docx —
the same library the practice generator uses — so a file Word would reject fails
here rather than in front of a building control officer.
"""
from playwright.sync_api import sync_playwright
import json, os, zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

# The practice this test signs the document as. It is set on the page, not read from a profile
# and not compiled into the app: the app must hard-code no practice at all, so the only way a
# name can reach the cover is by being supplied. A test that asserted a real firm would go green
# on exactly the white-label bug it is here to catch.
TEST_PRACTICE = "Marchmont Ridley Architects"
TEST_DESIGNER = "R Ridley"
TEST_LOGO = ("data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
             "AAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")
# Never on a generated document: the vendor's practice, or the software's own name.
FORBIDDEN = ["SY Design Studio", "Salman", "sydesignstudio", "Durham Avenue", "Specline"]
URL = "file://" + os.path.join(ROOT, "dist", "preview.html").replace("\\", "/")
DL = os.path.join(ROOT, "dist", "dl")
os.makedirs(DL, exist_ok=True)
R = []
def ok(n, c, x=""): R.append((("PASS" if c else "FAIL"), n, x))
def calc(pg,label):
    """open a calculator tab in the current category — the calculators sit behind tabs"""
    t=pg.locator('.stab:has-text("%s")' % label)
    if t.count(): t.first.click(); pg.wait_for_timeout(250)
    return t.count()

with sync_playwright() as p:
    b = p.chromium.launch()
    ctx = b.new_context(viewport={'width': 1500, 'height': 980}, accept_downloads=True)
    pg = ctx.new_page(); errs = []
    pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.goto(URL); pg.wait_for_timeout(700)
    pg.evaluate("try{localStorage.clear()}catch(e){}")
    pg.goto(URL); pg.wait_for_timeout(700)

    # a full extension job, everything selected, so every section is exercised
    pg.click('.tile[data-k="extension"]'); pg.wait_for_timeout(600)
    ok("W1 Word button present", pg.locator("#btnDocx").count() == 1)
    # The practice profile starts blank — nothing is compiled in — so give it one, the way a
    # subscriber fills it in on the Practice page.
    ok("W0 the app carries no practice of its own",
       pg.evaluate("PRACTICE_BLANK.name === '' && PRACTICE_BLANK.logo === '' && pName()")
       == "[Practice name]", str(pg.evaluate("[PRACTICE_BLANK.name, PRACTICE_BLANK.logo]")))
    pg.evaluate("""([n, d, l]) => { P.name=n; P.designer=d; P.addr='7 Fenwick Row, Leeds LS1 4AB';
        P.email='studio@marchmontridley.co.uk'; P.accent='#1F5C7A';
        P.logo=l; P.logoW=1; P.logoH=1; savePractice(); }""",
        [TEST_PRACTICE, TEST_DESIGNER, TEST_LOGO])
    pg.evaluate("""S.data.job='1150'; S.data.project='Single-storey rear extension';
        S.data.address='42 Hollybank Road, Hounslow'; S.data.client='Mr & Mrs Ahmed';
        S.sel=allBU().map((_,i)=>i); renderSteps(); renderStage(); renderPaper(); save();""")
    pg.wait_for_timeout(400)

    # add a calculated wall so section 4.0 has working in it
    pg.locator('button.step:has-text("External Walls")').click(); pg.wait_for_timeout(300)
    calc(pg,"Cavity wall calculator")
    if pg.locator("#addWall").count():
        pg.click("#addWall"); pg.wait_for_timeout(400)
    ok("W2 a calculated build-up is on the job", pg.evaluate("calcsOnJob().length") >= 1,
       str(pg.evaluate("calcsOnJob().length")))

    with pg.expect_download(timeout=40000) as di:
        pg.click("#btnDocx")
    d = di.value
    path = os.path.join(DL, d.suggested_filename)
    d.save_as(path)
    ok("W3 downloads with a .docx name", d.suggested_filename.endswith(".docx"), d.suggested_filename)
    size = os.path.getsize(path)
    ok("W4 file is a plausible size", size > 20000, f"{size} bytes")

    # a .docx is a zip: check the parts Word requires are all present
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        bad = z.testzip()
    ok("W5 zip is not corrupt", bad is None, str(bad))
    for part in ["[Content_Types].xml", "_rels/.rels", "word/document.xml",
                 "word/styles.xml", "word/_rels/document.xml.rels",
                 "word/header1.xml", "word/footer1.xml"]:
        ok(f"W6 part present: {part}", part in names, ", ".join(names[:9]))
    ok("W7 practice logo embedded", any(n.startswith("word/media/logo.") for n in names), ", ".join(names))

    # python-docx is the real test: it parses the XML the way Word does
    try:
        from docx import Document
        doc = Document(path)
        text = "\n".join(p.text for p in doc.paragraphs)
        ok("W8 python-docx opens the file", True)
        # The document must carry the practice that was supplied, and nobody else: not the
        # vendor's own practice, and not Specline, which is the software rather than the designer.
        ok("W9 cover carries the supplied practice",
           TEST_PRACTICE in text and TEST_DESIGNER in text,
           "missing " + TEST_PRACTICE)
        leaked = [f for f in FORBIDDEN if f in text]
        ok("W9b no vendor or product identity on the document", not leaked, ", ".join(leaked))
        ok("W10 responsibility statement present",
           "is the named designer and remains responsible" in text)
        ok("W11 compliance wording exact",
           "not an approval of it" in text and "building control body" in text)
        ok("W12 Part A and Part B headings present",
           "PART A — CONSTRUCTION BUILD-UPS" in text.upper() and "PART B" in text.upper())
        ok("W13 U-value working section present", "U-VALUE CALCULATIONS" in text.upper())
        ok("W14 a clause paragraph came through",
           "cavity" in text.lower() and len(text) > 20000, f"{len(text)} chars")
        # tables: schedule + cover + at least one working table
        ok("W15 tables present", len(doc.tables) >= 3, str(len(doc.tables)))
        sched = "\n".join(c.text for c in doc.tables[1].rows[0].cells) if len(doc.tables) > 1 else ""
        ok("W16 schedule table has the expected columns",
           "REF" in sched.upper() and "BUILD-UP" in sched.upper(), sched.replace("\n", " ")[:60])
        refs = "\n".join(c.text for r in doc.tables[1].rows for c in r.cells)
        ok("W17 build-up references numbered per job", "EW1" in refs and "GF1" in refs,
           refs.replace("\n", " ")[:80])
    except ImportError:
        ok("W8 python-docx opens the file", False, "python-docx not installed")

    # the superscript characters must survive; Word handles them, unlike the PDF fonts
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    ok("W18 m²K survives into the XML", "m²K" in xml or "W/m²K" in xml)
    ok("W19 no unescaped ampersands", "& " not in xml.replace("&amp; ", ""))
    # The accent is the practice's, darkened for text. SY Design Studio's orange was hard-coded
    # here until 7 September 2026, colouring every heading and reference on every practice's
    # specification; neither it nor Specline's petrol may appear on a document.
    ok("W21 headings take the practice's accent", "18485F" in xml.upper(),
       "expected the darkened #1F5C7A")
    for brand in ("B5640A", "E8850C", "F5900A", "0E6E85"):
        ok("W22 no fixed brand colour: " + brand, brand not in xml.upper())

    # W22b a draft download is marked for a practice that pays per specification. Without this
    # the Download button hands over the same document the Issue button charges for, so the
    # mark is the difference between what is free and what is paid for. The server decides who
    # is marked (drafts_are_marked in site/app/billing.php); here the flag is set directly,
    # because these suites run against dist/preview.html with no server behind them.
    with zipfile.ZipFile(path) as z:
        clean_hdr = z.read("word/header1.xml").decode("utf-8")
    ok("W22b an unmarked download carries no draft mark", "NOT FOR ISSUE" not in clean_hdr.upper())

    pg.evaluate("MARK_DRAFTS = true; renderPaper();")
    pg.wait_for_timeout(300)
    ok("W22c the preview says so, and says why",
       pg.locator("#paper .draftband").count() == 1
       and "credit is spent when you issue" in pg.locator("#paper .draftband").inner_text())

    with pg.expect_download(timeout=40000) as di2:
        pg.click("#btnDocx")
    d2 = di2.value
    p2 = os.path.join(DL, d2.suggested_filename)
    d2.save_as(p2)
    ok("W22d the draft file is named a draft", "_DRAFT" in d2.suggested_filename, d2.suggested_filename)
    with zipfile.ZipFile(p2) as z:
        hdr = z.read("word/header1.xml").decode("utf-8")
        body2 = z.read("word/document.xml").decode("utf-8")
    # The mark has to be in the HEADER, because that is what repeats on every page. In the body
    # it would sit on page one and be scrolled past.
    ok("W22e the mark is in the running header, so it is on every page",
       "NOT FOR ISSUE" in hdr.upper(), hdr[:160])
    ok("W22f and the specification itself is unchanged", "NOT FOR ISSUE" not in body2.upper())

    # Issuing is the paid act, so the document it produces is never marked.
    pg.evaluate("S.data.rev='P01'; S.history=[];")
    with pg.expect_download(timeout=40000) as di3:
        pg.evaluate("makeDoc('docx', true)")
    d3 = di3.value
    ok("W22g THE ISSUED DOCUMENT IS NEVER MARKED", "_DRAFT" not in d3.suggested_filename,
       d3.suggested_filename)
    p3 = os.path.join(DL, d3.suggested_filename)
    d3.save_as(p3)
    with zipfile.ZipFile(p3) as z:
        ok("W22h nor is its header", "NOT FOR ISSUE" not in z.read("word/header1.xml").decode("utf-8").upper())
    pg.evaluate("MARK_DRAFTS = false; renderPaper();")
    pg.wait_for_timeout(200)
    ok("W22i and with nothing to mark the band is gone", pg.locator("#paper .draftband").count() == 0)

    # W23 the practice page: three views, and each tab keeps the state the glance used to give
    pg.evaluate("go('practice')"); pg.wait_for_timeout(400)
    labs = [t.strip().splitlines()[0] for t in pg.locator("#practicepage .stab").all_inner_texts()]
    ok("W23 the practice page splits into identity, logo and plan",
       labs == ["Identity", "Logo", "Plan and seats"], str(labs))
    ok("W23 one view at a time", pg.locator("#practicepage .card:not([hidden])").count() == 1)
    ok("W23 the logo tab says which the cover carries",
       pg.locator('[data-stab="logo"] .stn').inner_text().strip() in ("Image", "Wordmark"),
       pg.locator('[data-stab="logo"] .stn').inner_text())
    ok("W23 a complete identity carries no warning",
       pg.locator('[data-stab="ident"] .stn').inner_text().strip() == "",
       pg.locator('[data-stab="ident"] .stn').inner_text())
    # clearing a field that prints on the cover must say so on the tab, at once
    pg.fill('#practicepage input[data-p="designer"]', ""); pg.wait_for_timeout(300)
    ok("W23 an emptied field is counted on the tab",
       "to fill" in pg.locator('[data-stab="ident"] .stn').inner_text(),
       pg.locator('[data-stab="ident"] .stn').inner_text().strip())
    pg.fill('#practicepage input[data-p="designer"]', TEST_DESIGNER); pg.wait_for_timeout(300)
    ok("W23 and uncounted when it is filled again",
       pg.locator('[data-stab="ident"] .stn').inner_text().strip() == "")
    # the seats tab counts against the plan, and says when it is over
    pg.locator('#practicepage .stab:has-text("Plan")').click(); pg.wait_for_timeout(300)
    ok("W23 the plan tab holds the seats", pg.locator("#practicepage .seats").count() == 1)
    pg.evaluate("P.plan='solo'; P.users=['a@x.co','b@x.co']; savePractice(); renderPractice();")
    pg.wait_for_timeout(300)
    chip = pg.locator('[data-stab="plan"] .stn')
    ok("W23 over the seat limit is named on the tab and marked",
       chip.inner_text().strip() == "2/1" and "bad" in (chip.get_attribute("class") or ""),
       chip.inner_text().strip() + " / " + str(chip.get_attribute("class")))
    ok("W23 the practice page still leaves for the desk",
       (pg.click("#practiceDone"), pg.wait_for_timeout(400), pg.evaluate("S.route") == "home")[2])

    ok("W20 no page errors", len(errs) == 0, "; ".join(errs[:2]))
    print(json.dumps([{"r": a, "t": b_, "x": c} for a, b_, c in R]))
    print("PAGE ERRORS:", errs[:4])
    b.close()
